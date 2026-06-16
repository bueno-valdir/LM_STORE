#!/usr/bin/env python3
"""
Gera a planilha-modelo (deploy/modelo-produtos.xlsx) que a cliente preenche.

Os campos de VALOR FIXO viram listas suspensas (dropdowns), pra nao errar:
  - categoria  -> nome amigavel das subcategorias (lido de src/config/categorias.ts)
  - tom        -> 0 a 10
  - destaque / promocao / disponivel -> sim / nao

Depois a cliente exporta como CSV e a importacao roda com scripts/importar-produtos.mjs
(o importador converte o nome amigavel da categoria de volta para o slug).

Rodar:  python3 scripts/gerar-planilha-modelo.py
"""

import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parent.parent
SAIDA = RAIZ / "deploy" / "modelo-produtos.xlsx"

# Slugs das categorias-mae (para separar das subcategorias ao ler o arquivo TS).
SLUGS_RAIZ = {"rosto", "olhos", "labios", "sobrancelhas", "skincare", "pinceis", "acessorios"}


def ler_subcategorias():
    """Le src/config/categorias.ts e retorna a lista de rotulos das subcategorias, em ordem."""
    txt = (RAIZ / "src" / "config" / "categorias.ts").read_text(encoding="utf-8")
    pares = re.findall(r"slug:\s*'([^']+)',\s*rotulo:\s*'([^']+)'", txt, re.DOTALL)
    return [rotulo for slug, rotulo in pares if slug not in SLUGS_RAIZ]


# Colunas: (cabecalho, nota explicativa, largura)
COLUNAS = [
    ("nome", "Nome do produto (obrigatorio).", 32),
    ("marca", "Marca. Escreva SEMPRE igual (ex.: Payot) para agrupar certo.", 16),
    ("categoria", "Escolha na lista suspensa.", 22),
    ("preco", "Preco normal. Ex.: 49.90 (use ponto ou virgula).", 12),
    ("preco_promocional", "Preco em promocao. Deixe vazio se nao houver.", 16),
    ("tom", "Nº de tons: 0 = sem tom; N = tons 1 a N. Lista suspensa.", 8),
    ("volume", "Volume/peso visivel ao cliente. Ex.: 30ml.", 12),
    ("descricao", "Descricao do produto.", 40),
    ("peso_g", "Peso para frete, em gramas. Ex.: 120.", 10),
    ("altura_cm", "Altura da embalagem (cm), para frete.", 10),
    ("largura_cm", "Largura da embalagem (cm), para frete.", 10),
    ("comprimento_cm", "Comprimento da embalagem (cm), para frete.", 13),
    ("estoque", "Unidades (produtos SEM tom). Vazio = nao controla.", 10),
    ("estoque_tons", "Unidades por tom separadas por | . Ex.: 5|0|2 (Tom1|Tom2|Tom3).", 16),
    ("destaque", "Aparece na home? Lista suspensa (sim/nao).", 10),
    ("promocao", "Esta em promocao? Lista suspensa (sim/nao).", 10),
    ("disponivel", "A venda? Lista suspensa (sim/nao).", 11),
    ("imagens", "URLs das fotos (links publicos) separadas por | . A 1a vira capa.", 50),
]

# Linhas de exemplo (ajudam a cliente a entender o formato). Serao apagadas por ela.
EXEMPLOS = [
    ["Base Liquida Matte Alta Cobertura", "Payot", "Base", 49.90, "", 3, "30ml",
     "Base liquida matte com alta cobertura, disfarca imperfeicoes e uniformiza o tom.",
     120, 12, 4, 4, "", "5|0|2", "sim", "nao", "sim",
     "https://exemplo.com/base-frente.jpg|https://exemplo.com/base-verso.jpg"],
    ["Serum Facial Vitamina C", "Eudora", "Sérum / Tratamentos", 79.90, 69.90, 0, "30ml",
     "Serum antioxidante com vitamina C para vico e uniformizacao.",
     110, 10, 4, 4, 8, "", "nao", "sim", "sim",
     "https://exemplo.com/serum.jpg"],
]

LINHAS_PREENCHIVEIS = 600  # ate onde os dropdowns valem


def main():
    rotulos = ler_subcategorias()
    wb = Workbook()

    # ---- aba Produtos ----
    ws = wb.active
    ws.title = "Produtos"

    cab_fill = PatternFill("solid", fgColor="7A4DBA")
    cab_font = Font(color="FFFFFF", bold=True)
    for i, (nome, nota, larg) in enumerate(COLUNAS, start=1):
        c = ws.cell(row=1, column=i, value=nome)
        c.fill = cab_fill
        c.font = cab_font
        c.alignment = Alignment(vertical="center")
        c.comment = Comment(nota, "LM Store")
        ws.column_dimensions[get_column_letter(i)].width = larg
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "A2"

    for r, linha in enumerate(EXEMPLOS, start=2):
        for cidx, val in enumerate(linha, start=1):
            ws.cell(row=r, column=cidx, value=val)

    # ---- aba Listas (fonte dos dropdowns de categoria) ----
    wl = wb.create_sheet("Listas")
    wl["A1"] = "categorias"
    for i, rot in enumerate(rotulos, start=2):
        wl.cell(row=i, column=1, value=rot)
    wl.column_dimensions["A"].width = 32
    wl.sheet_state = "hidden"

    def idx(nome):
        return next(i for i, c in enumerate(COLUNAS, start=1) if c[0] == nome)

    fim = LINHAS_PREENCHIVEIS
    # categoria -> referencia a aba Listas
    n_cat = len(rotulos) + 1
    dv_cat = DataValidation(
        type="list", formula1=f"=Listas!$A$2:$A${n_cat}", allow_blank=True,
        showErrorMessage=True, errorTitle="Categoria invalida",
        error="Escolha uma categoria da lista.")
    ws.add_data_validation(dv_cat)
    col = get_column_letter(idx("categoria"))
    dv_cat.add(f"{col}2:{col}{fim}")

    # tom -> 0..10
    dv_tom = DataValidation(type="list", formula1='"0,1,2,3,4,5,6,7,8,9,10"', allow_blank=True)
    ws.add_data_validation(dv_tom)
    col = get_column_letter(idx("tom"))
    dv_tom.add(f"{col}2:{col}{fim}")

    # sim/nao -> destaque, promocao, disponivel
    for campo in ("destaque", "promocao", "disponivel"):
        dv = DataValidation(type="list", formula1='"sim,nao"', allow_blank=True)
        ws.add_data_validation(dv)
        col = get_column_letter(idx(campo))
        dv.add(f"{col}2:{col}{fim}")

    # ---- aba Instrucoes ----
    wi = wb.create_sheet("Instrucoes")
    wi.column_dimensions["A"].width = 100
    instrucoes = [
        ("COMO PREENCHER", True),
        ("", False),
        ("1) Cada LINHA da aba 'Produtos' e um produto.", False),
        ("2) Campos com seta (lista suspensa): categoria, tom, destaque, promocao, disponivel.", False),
        ("   Clique na celula e escolha uma opcao. Nao digite a mao nesses campos.", False),
        ("3) tom: 0 = produto sem tom. N = produto com tons de 1 ate N (cliente escolhe no site).", False),
        ("4) estoque (sem tom) OU estoque_tons (com tom): unidades. Vazio = nao controla estoque.", False),
        ("   estoque_tons usa | entre os tons. Ex.: 5|0|2 = Tom1 tem 5, Tom2 esgotado, Tom3 tem 2.", False),
        ("5) imagens: cole os LINKS das fotos (precisam abrir no navegador), separados por | .", False),
        ("   A primeira foto vira a capa do produto.", False),
        ("6) preco: use ponto ou virgula (49.90 ou 49,90). preco_promocional so se houver promo.", False),
        ("7) As 2 primeiras linhas da aba Produtos sao EXEMPLOS. Apague antes de importar.", False),
        ("", False),
        ("DEPOIS DE PREENCHER", True),
        ("Arquivo > Fazer download > CSV (no Google Sheets) e envie o CSV para a importacao.", False),
        ("A importacao converte o nome da categoria para o codigo interno automaticamente.", False),
    ]
    for r, (txt, neg) in enumerate(instrucoes, start=1):
        c = wi.cell(row=r, column=1, value=txt)
        if neg:
            c.font = Font(bold=True, size=12, color="7A4DBA")

    wb.save(SAIDA)
    print(f"Gerado: {SAIDA}  ({len(rotulos)} categorias na lista)")


if __name__ == "__main__":
    main()
